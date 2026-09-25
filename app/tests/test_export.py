import unittest
from datetime import datetime, timezone
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock, MagicMock

from aiogram.fsm.context import FSMContext
from aiogram.fsm.storage.base import StorageKey
from aiogram.fsm.storage.memory import MemoryStorage
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.orm import Session

from app.handlers.export import download_excel
from app.models import Base, Employee, Operation, Product, ProductComponent
from app.repositories.export import ExportReadError, ExportRepository
from app.services.export import ExportService


class ExportTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self) -> None:
        self.engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(self.engine)
        self.db = Session(self.engine)
        session = AsyncMock()
        session.__aenter__.return_value = session
        session.execute.side_effect = self.db.execute
        self.session = session
        self.repository = ExportRepository(MagicMock(return_value=session))
        self.service = ExportService(self.repository)
        self.storage = MemoryStorage()
        self.state = FSMContext(
            self.storage, StorageKey(bot_id=1, chat_id=1, user_id=1)
        )

    async def asyncTearDown(self) -> None:
        self.db.close()
        self.engine.dispose()
        await self.storage.close()

    def seed(self) -> None:
        self.db.add(Employee(id=1, telegram_id="00123", name="Анна"))
        self.db.add(Product(
            id=1, article="A-01", name="Контроллер", price=Decimal("125.50"),
            quantity=7, comment="Тестовый товар",
        ))
        self.db.flush()
        self.db.add_all([
            Operation(
                id=1, product_id=1, op_type="Поступление", quantity=10,
                comment="Приёмка",
                created_at=datetime(2026, 1, 10, 12, 0, tzinfo=timezone.utc),
            ),
            Operation(
                id=2, product_id=1, op_type="Отгрузка", quantity=3,
                comment="Выдача",
                created_at=datetime(2026, 1, 11, 12, 0, tzinfo=timezone.utc),
            ),
        ])
        self.db.commit()

    async def test_excel_contains_all_sheets_and_data(self) -> None:
        self.seed()
        content = await self.service.create_excel()
        workbook = load_workbook(BytesIO(content), data_only=True)

        self.assertEqual(
            workbook.sheetnames,
            ["Товары", "Сотрудники", "Движения", "Продукты"],
        )
        self.assertEqual(workbook["Товары"]["B2"].value, "A-01")
        self.assertEqual(workbook["Товары"]["E2"].value, 7)
        self.assertEqual(workbook["Сотрудники"]["B2"].value, "00123")
        self.assertEqual(list(workbook["Движения"].values), [
            ("Дата", "Наименование", "Поступление", "Отгрузка", "Комментарий"),
            (datetime(2026, 1, 11, 15), "Контроллер", 0, 3, "Выдача"),
            (datetime(2026, 1, 10, 15), "Контроллер", 10, 0, "Приёмка"),
        ])
        self.assertEqual(workbook["Движения"]["A2"].number_format, "dd.mm.yyyy hh:mm")

    async def test_export_contains_component_responsible_and_composition(self) -> None:
        self.seed()
        self.db.get(Product, 1).responsible = "Анна"
        self.db.add(Product(id=3, article="0018", name="Коробка", price=10, comment="-"))
        self.db.flush()
        self.db.add(Product(
            id=2, article="FIN", name="Прибор", kind="product", price=0, comment="-",
            components=[ProductComponent(component_id=1, quantity=3),
                        ProductComponent(component_id=3, quantity=2)],
        ))
        self.db.commit()
        workbook = load_workbook(BytesIO(await self.service.create_excel()))
        self.assertEqual(workbook["Товары"]["G2"].value, "Комплектующее")
        self.assertEqual(workbook["Товары"]["H2"].value, "Анна")
        self.assertEqual(workbook["Товары"]["G3"].value, "Продукт")
        self.assertEqual(list(workbook["Продукты"].values), [
            ("Артикул", "Наименование", "Состав"),
            ("FIN", "Прибор", "A-01 (3), 0018 (2)"),
        ])
        self.assertNotIn("Ответственный", list(workbook["Движения"].values)[0])

    async def test_text_is_not_interpreted_as_excel_formula(self) -> None:
        self.seed()
        product = self.db.get(Product, 1)
        product.name = "=1+1"
        product.comment = "#N/A"
        self.db.commit()
        workbook = load_workbook(BytesIO(await self.service.create_excel()))
        self.assertEqual(workbook["Товары"]["C2"].value, "=1+1")
        self.assertEqual(workbook["Товары"]["C2"].data_type, "s")
        self.assertEqual(workbook["Товары"]["F2"].value, "#N/A")
        self.assertEqual(workbook["Товары"]["F2"].data_type, "s")

    async def test_handler_sends_file_and_clears_state(self) -> None:
        service = AsyncMock(spec=ExportService)
        service.create_excel.return_value = b"excel"
        message = SimpleNamespace(answer=AsyncMock(), answer_document=AsyncMock())
        await self.state.set_state("unfinished")

        await download_excel(message, self.state, service)

        document = message.answer_document.call_args.args[0]
        self.assertTrue(document.filename.startswith("warehouse_"))
        self.assertTrue(document.filename.endswith(".xlsx"))
        self.assertEqual(document.data, b"excel")
        self.assertIsNone(await self.state.get_state())

    async def test_database_error_is_reported(self) -> None:
        self.session.execute.side_effect = SQLAlchemyError("Unavailable")
        with self.assertRaises(ExportReadError):
            await self.repository.get_all()

        service = AsyncMock(spec=ExportService)
        service.create_excel.side_effect = ExportReadError()
        message = SimpleNamespace(answer=AsyncMock(), answer_document=AsyncMock())
        with self.assertLogs("app.handlers.export", level="ERROR"):
            await download_excel(message, self.state, service)
        self.assertIn("Не удалось", message.answer.call_args.args[0])
        message.answer_document.assert_not_awaited()
