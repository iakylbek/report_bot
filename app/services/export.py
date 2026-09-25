from datetime import datetime, timezone
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Font

from app.repositories.export import ExportRepository


MOSCOW_TIMEZONE = ZoneInfo("Europe/Moscow")


class ExportService:
    def __init__(self, repository: ExportRepository) -> None:
        self.repository = repository

    async def create_excel(self) -> bytes:
        data = await self.repository.get_all()
        workbook = Workbook()
        workbook.remove(workbook.active)

        self._add_sheet(
            workbook,
            "Товары",
            ["ID", "Артикул", "Название", "Цена", "Количество", "Комментарий", "Тип", "Ответственный"],
            [
                [
                    row["id"], row["article"], row["name"], row["price"],
                    row["quantity"], row["comment"],
                    "Продукт" if row["kind"] == "product" else "Комплектующее", row["responsible"],
                ]
                for row in data["products"]
            ],
        )
        self._add_sheet(
            workbook,
            "Сотрудники",
            ["ID", "Telegram ID", "Имя"],
            [
                [row["id"], row["telegram_id"], row["name"]]
                for row in data["employees"]
            ],
        )

        self._add_sheet(
            workbook,
            "Движения",
            ["Дата", "Наименование", "Поступление", "Отгрузка", "Комментарий"],
            [
                [
                    self._moscow_time(row["created_at"]), row["name"],
                    row["quantity"] if row["op_type"] == "Поступление" else 0,
                    row["quantity"] if row["op_type"] == "Отгрузка" else 0,
                    row["comment"],
                ]
                for row in data["operations"]
            ],
        )

        compositions: dict[int, list[str]] = {}
        for row in data["components"]:
            compositions.setdefault(row["product_id"], []).append(
                f"{row['component_article']} ({row['quantity']})"
            )
        self._add_sheet(
            workbook, "Продукты",
            ["Артикул", "Наименование", "Состав"],
            [
                [row["article"], row["name"], ", ".join(compositions.get(row["id"], []))]
                for row in data["products"] if row["kind"] == "product"
            ],
        )
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    @staticmethod
    def _moscow_time(value: datetime) -> datetime:
        if value.tzinfo is None:
            value = value.replace(tzinfo=timezone.utc)
        # Excel хранит даты без часового пояса; предварительно переводим в МСК.
        return value.astimezone(MOSCOW_TIMEZONE).replace(tzinfo=None)

    @staticmethod
    def _add_sheet(
        workbook: Workbook, title: str, headers: list[str], rows: list[list]
    ) -> None:
        sheet = workbook.create_sheet(title)
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row in rows:
            sheet.append(row)
            # Названия и комментарии должны оставаться текстом, даже начинаясь с «=».
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                elif isinstance(cell.value, datetime):
                    cell.number_format = "dd.mm.yyyy hh:mm"
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            width = min(max(len(str(cell.value or "")) for cell in column) + 2, 45)
            sheet.column_dimensions[column[0].column_letter].width = width
